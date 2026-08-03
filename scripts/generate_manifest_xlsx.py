#!/usr/bin/env python3
"""Generate an .xlsx workbook per manifest type, pairing each CSV template in
manifest_templates/ with its field documentation in docs/. Each workbook has
two visible tabs: "Data Dictionary" (from the docs/*.md table) and
"Manifest Template" (the CSV header row, with dropdowns for any column that
has an Allowed Values list in the data dictionary).
"""

import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = REPO_ROOT / "docs"
TEMPLATES_DIR = REPO_ROOT / "manifest_templates"
OUT_DIR = REPO_ROOT / "excel_templates"

DICT_HEADERS = [
    "Column Name",
    "Required",
    "Explanation",
    "Allowed Values",
    "Data Type",
    "Example Entry",
    "BTI Derived",
]

# Column widths for the Data Dictionary sheet, keyed by header.
DICT_COLUMN_WIDTHS = {
    "Column Name": 28,
    "Required": 10,
    "Explanation": 60,
    "Allowed Values": 45,
    "Data Type": 12,
    "Example Entry": 22,
    "BTI Derived": 13,
}

TEMPLATE_ROWS_TO_VALIDATE = 1000
HEADER_FONT = Font(bold=True)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def parse_data_dictionary(md_path: Path) -> list[list[str]]:
    """Pull the pipe-delimited field table out of a docs/*.md file."""
    lines = md_path.read_text().splitlines()

    header_idx = next(
        i for i, line in enumerate(lines) if line.strip().startswith("| Column Name")
    )
    rows = []
    for line in lines[header_idx + 2 :]:
        if not line.strip().startswith("|"):
            break
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def read_csv_header(csv_path: Path) -> list[str]:
    with csv_path.open(newline="") as fh:
        return next(csv.reader(fh))


def allowed_values_for(dict_rows: list[list[str]]) -> dict[str, list[str]]:
    """Map column name -> list of allowed values parsed from quoted strings."""
    values_by_column = {}
    for row in dict_rows:
        column_name, allowed_values = row[0], row[3]
        found = re.findall(r'"([^"]*)"', allowed_values)
        if found:
            values_by_column[column_name] = found
    return values_by_column


def build_data_dictionary_sheet(wb: Workbook, dict_rows: list[list[str]]) -> None:
    ws = wb.create_sheet("Data Dictionary")
    ws.append(DICT_HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for row in dict_rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
    for idx, header in enumerate(DICT_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = DICT_COLUMN_WIDTHS[header]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_manifest_template_sheet(
    wb: Workbook, csv_header: list[str], values_by_column: dict[str, list[str]]
) -> None:
    ws = wb.create_sheet("Manifest Template")
    ws.append(csv_header)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for idx in range(1, len(csv_header) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 24
    ws.freeze_panes = "A2"

    lists_ws = None
    for col_idx, column_name in enumerate(csv_header, start=1):
        values = values_by_column.get(column_name)
        if not values:
            continue
        if lists_ws is None:
            lists_ws = wb.create_sheet("Lists")
            lists_ws.sheet_state = "hidden"
        list_col = lists_ws.max_column + 1 if lists_ws.max_column > 1 or lists_ws["A1"].value else 1
        list_col_letter = get_column_letter(list_col)
        lists_ws.cell(row=1, column=list_col, value=column_name)
        for i, value in enumerate(values, start=2):
            lists_ws.cell(row=i, column=list_col, value=value)

        dv = DataValidation(
            type="list",
            formula1=f"'Lists'!${list_col_letter}$2:${list_col_letter}${len(values) + 1}",
            allow_blank=True,
            showDropDown=False,
        )
        ws.add_data_validation(dv)
        col_letter = get_column_letter(col_idx)
        dv.add(f"{col_letter}2:{col_letter}{TEMPLATE_ROWS_TO_VALIDATE}")


def generate_workbook(csv_path: Path, md_path: Path, out_path: Path) -> None:
    dict_rows = parse_data_dictionary(md_path)
    csv_header = read_csv_header(csv_path)
    values_by_column = allowed_values_for(dict_rows)

    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet
    build_data_dictionary_sheet(wb, dict_rows)
    build_manifest_template_sheet(wb, csv_header, values_by_column)
    wb.active = 0

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    csv_paths = sorted(TEMPLATES_DIR.glob("*_manifest_template.csv"))
    if not csv_paths:
        raise SystemExit(f"No manifest templates found in {TEMPLATES_DIR}")

    for csv_path in csv_paths:
        stem = csv_path.stem
        md_path = DOCS_DIR / f"{stem}.md"
        if not md_path.exists():
            raise SystemExit(f"Missing data dictionary doc for {csv_path.name}: {md_path}")
        out_path = OUT_DIR / f"{stem}.xlsx"
        generate_workbook(csv_path, md_path, out_path)
        print(f"Wrote {out_path.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
